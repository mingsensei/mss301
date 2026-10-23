import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_slot13_docx():
    base_docx = 'Slot13/clean_base.docx'
    output_docx = 'Slot13/Slot13_Postman_Test_Report.docx'

    doc = docx.Document(base_docx)

    # 1. Update Header & Footer
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            p.text = ""
        hp = sec.header.paragraphs[0]
        hp.style = 'Header'
        hrun = hp.add_run('Báo cáo Thực hành Slot 13 – Saga Pattern: Order Saga (Orchestration) – MSS301')
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(9.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Update footer runs font to Times New Roman
        for fp in sec.footer.paragraphs:
            for fr in fp.runs:
                fr.font.name = 'Times New Roman'

    # 2. Clear Document Body
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(doc.tables):
        t._element.getparent().remove(t._element)

    # Helper styling functions
    def set_cell_border(cell, **kwargs):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge_data = kwargs.get(edge)
            if edge_data:
                tag = f'w:{edge}'
                element = parse_xml(f'<{tag} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="{edge_data.get("space", "0")}" w:color="{edge_data.get("color", "auto")}"/>')
                tcBorders.append(element)
            else:
                tag = f'w:{edge}'
                element = parse_xml(f'<{tag} {nsdecls("w")} w:val="none"/>')
                tcBorders.append(element)
        tcPr.append(tcBorders)

    def set_cell_shading(cell, color_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.style = '1.Tieude1'
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.style = '1.Tieude2'
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.style = '0.NoiDung'
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25

        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Times New Roman"
            rb.font.size = Pt(12)
            rb.font.bold = True
            rb.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = False
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.style = '0.NoiDung'
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.25

        rbullet = p.add_run("• ")
        rbullet.font.name = "Times New Roman"
        rbullet.font.size = Pt(12)
        rbullet.font.bold = True
        rbullet.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Times New Roman"
            rb.font.size = Pt(12)
            rb.font.bold = True
            rb.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = False
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    def add_image_space(img_num, img_title, explanation_text):
        """
        Leaves blank space (empty centered paragraph) for user to paste image,
        followed by Caption and Explanation in Style Tieudehinh.
        """
        # Blank space paragraph for pasting screenshot
        p_space = doc.add_paragraph()
        p_space.style = 'Normal'
        p_space.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_space.paragraph_format.space_before = Pt(12)
        p_space.paragraph_format.space_after = Pt(12)
        r_sp = p_space.add_run("\n\n")
        r_sp.font.name = "Times New Roman"

        # Caption (Tieudehinh)
        p_cap = doc.add_paragraph()
        p_cap.style = 'Tieudehinh'
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True

        rc = p_cap.add_run(f"Hình {img_num}: {img_title}")
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(11)
        rc.font.bold = True
        rc.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Explanation (Tieudehinh)
        p_exp = doc.add_paragraph()
        p_exp.style = 'Tieudehinh'
        p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_exp.paragraph_format.space_before = Pt(2)
        p_exp.paragraph_format.space_after = Pt(8)

        re_lbl = p_exp.add_run("Giải thích kết quả: ")
        re_lbl.font.name = "Times New Roman"
        re_lbl.font.size = Pt(10)
        re_lbl.font.bold = True
        re_lbl.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        re_val = p_exp.add_run(explanation_text)
        re_val.font.name = "Times New Roman"
        re_val.font.size = Pt(10)
        re_val.font.bold = False
        re_val.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    # 3. Document Title Block
    p_title = doc.add_paragraph()
    p_title.style = 'Normal'
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)

    rt1 = p_title.add_run("BÁO CÁO THỰC HÀNH SLOT 13 – MSS301\n")
    rt1.font.name = "Times New Roman"
    rt1.font.size = Pt(20)
    rt1.font.bold = True
    rt1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    rt2 = p_title.add_run("SAGA PATTERN: ORDER SAGA (ORCHESTRATION)\n")
    rt2.font.name = "Times New Roman"
    rt2.font.size = Pt(14)
    rt2.font.bold = True
    rt2.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    rt3 = p_title.add_run("KIỂM THỬ GIAO DỊCH PHÂN TÁN MICROSERVICES VỚI POSTMAN")
    rt3.font.name = "Times New Roman"
    rt3.font.size = Pt(13)
    rt3.font.bold = True
    rt3.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # 4. Info Table
    info_table = doc.add_table(rows=7, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    col_widths = [Inches(2.2), Inches(3.8)]
    for row in info_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    info_data = [
        ("Môn học:", "MSS301 – Kiến trúc Hướng Dịch vụ & Microservices"),
        ("Bài thực hành:", "Slot 13 – Saga Pattern: Order Saga (Orchestration)"),
        ("Giảng viên hướng dẫn:", "traltb@fe.edu.vn"),
        ("Sinh viên thực hiện:", "Dương Hồng Minh (MSSV: DE190220)"),
        ("Kiến trúc hệ thống:", "Microservices Polyglot (MySQL 8.4, MongoDB 8.0, SQL Server 2022)"),
        ("Phương pháp kiểm thử:", "Postman Collection Runner & Newman Automated CLI"),
        ("Ngày hoàn thành:", "23/10/2026")
    ]
    b_spec = dict(val="single", sz="4", space="0", color="D1D5DB")
    for i, (label, val) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_margins(c0, top=90, bottom=90, left=130, right=130)
        set_cell_shading(c0, "F8FAFC")
        set_cell_border(c0, top=b_spec, bottom=b_spec, left=b_spec, right=b_spec)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before, p0.paragraph_format.space_after = Pt(2), Pt(2)
        r0 = p0.add_run(label)
        r0.font.name, r0.font.size, r0.font.bold = "Times New Roman", Pt(11), True
        r0.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        set_cell_margins(c1, top=90, bottom=90, left=130, right=130)
        set_cell_shading(c1, "FFFFFF")
        set_cell_border(c1, top=b_spec, bottom=b_spec, left=b_spec, right=b_spec)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before, p1.paragraph_format.space_after = Pt(2), Pt(2)
        r1 = p1.add_run(val)
        r1.font.name, r1.font.size = "Times New Roman", Pt(11)
        r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # 5. TOC
    sp = doc.add_paragraph()
    sp.style = 'Normal'
    sp.paragraph_format.space_before, sp.paragraph_format.space_after = Pt(6), Pt(6)

    p_toc_head = doc.add_paragraph()
    p_toc_head.style = 'TOC Heading'
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_head.paragraph_format.space_before = Pt(16)
    p_toc_head.paragraph_format.space_after = Pt(10)
    rtoc = p_toc_head.add_run("MỤC LỤC CHI TIẾT")
    rtoc.font.name, rtoc.font.size, rtoc.font.bold = "Times New Roman", Pt(16), True

    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_before, p_toc.paragraph_format.space_after = Pt(0), Pt(10)
    p_toc_el = p_toc._p
    p_toc_el.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/></w:rPr><w:fldChar w:fldCharType="begin"/></w:r>'))
    p_toc_el.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/></w:rPr><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'))
    p_toc_el.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/></w:rPr><w:fldChar w:fldCharType="separate"/></w:r>'))
    p_toc_el.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/></w:rPr><w:fldChar w:fldCharType="end"/></w:r>'))

    p_br = doc.add_paragraph()
    p_br.paragraph_format.space_before, p_br.paragraph_format.space_after = Pt(0), Pt(0)
    p_br.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)

    # =========================================================================
    # SECTION 1: KHỞI CHẠY HẠ TẦNG VÀ KIỂM TRA HỆ THỐNG
    # =========================================================================
    add_heading_1("1. KHỞI CHẠY HẠ TẦNG DOCKER COMPOSE VÀ KIỂM TRA HỆ THỐNG")
    add_body("Hạ tầng kiểm thử bao gồm 3 cơ sở dữ liệu phân tán (MySQL 8.4, MongoDB 8.0, SQL Server 2022) và 3 microservices Spring Boot 4.1.1 (order-service, inventory-service, payment-service) được container hóa bằng Docker Compose.")

    add_heading_2("1.1. Khởi động 3 Database và 3 Microservices bằng Docker Compose")
    add_body("Thực thi lệnh 'docker compose up -d' để khởi động toàn bộ các container dịch vụ và cơ sở dữ liệu, sau đó kiểm tra trạng thái sức khỏe bằng 'docker compose ps':")

    add_image_space(
        1,
        "Trạng thái 6 containers hoạt động thành công trong Docker Compose",
        "Xác nhận 3 cơ sở dữ liệu (MySQL 8.4 - port 3307, MongoDB 8.0 - port 27017, SQL Server 2022 - port 14330) và 3 microservices (ports 8081, 8082, 8083) đều hoạt động bình thường, sẵn sàng phục vụ các kịch bản kiểm thử."
    )

    add_heading_2("1.2. Kiểm tra dữ liệu khởi tạo ban đầu trực tiếp trên 3 Database")
    add_body("Kiểm tra trạng thái dữ liệu ban đầu trực tiếp từ terminal vào từng container database (product P1 có 10, ví C1 có 5000, ví C2 có 100, danh sách đơn hàng rỗng):")

    add_image_space(
        2,
        "Kết quả kiểm tra dữ liệu khởi tạo ban đầu trực tiếp trên 3 Database",
        "Xác nhận trạng thái ban đầu của hệ thống: P1 có tồn kho 10, C1 có 5000.0 (đủ điều kiện cho Happy Path), C2 có 100.0 (sẽ kích hoạt lỗi không đủ tiền ở KB2)."
    )

    # =========================================================================
    # SECTION 2: SMOKE TEST KIỂM TRA SẴN SÀNG CỦA CÁC SERVICE
    # =========================================================================
    add_heading_1("2. KIỂM TRA SẴN SÀNG CỦA CÁC DỊCH VỤ (SMOKE TEST POSTMAN)")
    add_body("Thực thi các request trong thư mục '0 - Kiểm tra service' trên Postman để xác nhận cả 3 microservices đều phản hồi ổn định qua giao thức HTTP REST (GET /api/inventory/products, GET /api/payments/wallets, GET /api/orders):")

    add_image_space(
        3,
        "Kết quả thực thi 3 request Smoke Test ban đầu trong Postman",
        "Cả 3 microservices phản hồi mã HTTP 200 OK, xác nhận tầng Application sẵn sàng tiếp nhận các giao dịch phân tán."
    )

    # =========================================================================
    # SECTION 3: KB1 - HAPPY PATH
    # =========================================================================
    add_heading_1("3. KỊCH BẢN 1 (KB1) – HAPPY PATH: ĐẶT HÀNG VÀ THANH TOÁN THÀNH CÔNG")
    add_body("Mô tả kịch bản: Khách hàng C1 đặt mua 2 sản phẩm P1 với tổng tiền 2000.0. Kho MongoDB còn 10 P1, ví SQL Server C1 có 5000.0. Luồng Saga điều phối tuần tự thành công toàn bộ chuỗi transaction, cập nhật đơn hàng thành CONFIRMED và Saga COMPLETED.")

    add_heading_2("3.1. [KB1] Bước 1 – Snapshot tồn kho P1 và số dư ví C1 ban đầu")
    add_body("Gửi request GET /api/inventory/products và GET /api/payments/wallets để lưu lại số liệu tồn kho P1 (10) và số dư ví C1 (5000.0) vào biến Collection Variables trước khi đặt hàng:")

    add_image_space(
        4,
        "[Postman KB1] Snapshot trạng thái tồn kho P1 (10) và số dư ví C1 (5000)",
        "Postman ghi nhận thành công trạng thái snapshot: stock_P1 = 10, bal_C1 = 5000.0 làm căn cứ đối soát sau khi đặt hàng."
    )

    add_heading_2("3.2. [KB1] Bước 2 – Gửi yêu cầu đặt hàng POST /api/orders")
    add_body("Gửi request POST http://localhost:8081/api/orders với body {customerId: \"C1\", productId: \"P1\", quantity: 2, amount: 2000.0}:")

    add_image_space(
        5,
        "[Postman KB1] Request đặt hàng thành công đơn hàng số 1 (CONFIRMED, COMPLETED)",
        "Saga Orchestrator đã thực thi thành công toàn bộ chuỗi transaction: Order được xác nhận CONFIRMED và Saga hoàn tất COMPLETED mà không xảy ra bất kỳ lỗi ngoại lệ nào."
    )

    add_heading_2("3.3. [KB1] Bước 3 – Kiểm tra đối soát biến động kho và ví sau khi đặt hàng")
    add_body("Gửi request kiểm tra kho và ví để đối soát độ chênh lệch: tồn kho P1 giảm từ 10 xuống 8 (-2), số dư ví C1 giảm từ 5000 xuống 3000 (-2000):")

    add_image_space(
        6,
        "[Postman KB1] Kiểm tra đối soát tồn kho P1 còn 8 và số dư C1 còn 3000",
        "Hệ thống đã trừ chính xác 2 sản phẩm P1 trên MongoDB và trừ chính xác 2000.0 tiền trong ví C1 trên SQL Server, chứng minh tính toàn vẹn dữ liệu đa hệ quản trị."
    )

    add_heading_2("3.4. [KB1] Bước 4 – Kiểm tra chuỗi nhật ký điều phối Saga Logs")
    add_body("Truy vấn lịch sử điều phối của đơn hàng số 1 qua API GET http://localhost:8081/api/orders/1/saga-logs (4 bước CREATE_ORDER, RESERVE_INVENTORY, PROCESS_PAYMENT, CONFIRM_ORDER đều đạt kết quả SUCCESS):")

    add_image_space(
        7,
        "[Postman KB1] Chuỗi nhật ký thực thi Saga Logs của đơn hàng KB1 (4 bước SUCCESS)",
        "Nhật ký Saga Logs minh chứng tiến trình điều phối tuần tự của Orchestrator: mọi bước đều thành công trọn vẹn theo chiều xuôi."
    )

    add_heading_2("3.5. [KB1] Bước 5 – Xác minh dữ liệu trực tiếp trên 3 Database")
    add_body("Truy vấn trực tiếp MySQL (orders id=1 CONFIRMED), MongoDB (product P1 stock=8, reservation RESERVED) và SQL Server (wallet C1 balance=3000, payment PAID):")

    add_image_space(
        8,
        "Kết quả truy vấn đối soát dữ liệu thực tế trên 3 Database sau KB1",
        "Dữ liệu thực tế tại cả 3 database độc lập hoàn toàn đồng nhất, chứng minh giao dịch phân tán Happy Path thành công tuyệt đối."
    )

    # =========================================================================
    # SECTION 4: KB2 - COMPENSATING TRANSACTION
    # =========================================================================
    add_heading_1("4. KỊCH BẢN 2 (KB2) – THANH TOÁN THẤT BẠI VÀ BÙ TRỪ PHÂN TÁN (COMPENSATING TRANSACTION)")
    add_body("Mô tả kịch bản: Khách hàng C2 (chỉ có 100.0 trong ví) đặt mua 1 sản phẩm P1 trị giá 1000.0. Sau khi inventory-service giữ hàng thành công, payment-service trả về lỗi HTTP 402 Payment Required. Orchestrator tự động kích hoạt Compensating Transaction để giải phóng hàng trong MongoDB và hủy đơn hàng trong MySQL.")

    add_heading_2("4.1. [KB2] Bước 1 – Snapshot tồn kho P1 và số dư ví C2 trước khi chạy KB2")
    add_body("Gửi request GET /api/inventory/products và GET /api/payments/wallets ghi nhận trạng thái mốc: tồn kho P1 là 8 và số dư ví C2 là 100.0:")

    add_image_space(
        9,
        "[Postman KB2] Snapshot trạng thái tồn kho P1 (8) và số dư ví C2 (100)",
        "Hệ thống ghi nhận điểm mốc snapshot trước KB2: kho P1 còn 8, ví C2 chỉ có 100.0, không đủ khả năng chi trả đơn hàng 1000.0."
    )

    add_heading_2("4.2. [KB2] Bước 2 – Gửi yêu cầu đặt hàng gây lỗi thanh toán POST /api/orders")
    add_body("Gửi request POST http://localhost:8081/api/orders với body {customerId: \"C2\", productId: \"P1\", quantity: 1, amount: 1000.0} (kỳ vọng status CANCELLED, sagaStatus COMPENSATED, lỗi 402):")

    add_image_space(
        10,
        "[Postman KB2] Request đặt hàng kích hoạt lỗi 402 và cơ chế bù trừ (CANCELLED, COMPENSATED)",
        "Orchestrator bắt được ngoại lệ 402 từ payment-service, chuyển trạng thái Saga sang COMPENSATING, thực thi release kho và cập nhật đơn hàng thành CANCELLED / COMPENSATED."
    )

    add_heading_2("4.3. [KB2] Bước 3 – Kiểm tra đối soát bảo toàn kho P1 và số dư ví C2 sau bù trừ")
    add_body("Gửi request kiểm tra kho và ví để xác minh tính toàn vẹn: tồn kho P1 vẫn nguyên vẹn 8 (bị trừ tạm còn 7 rồi được hoàn trả về 8), số dư ví C2 vẫn giữ nguyên 100.0:")

    add_image_space(
        11,
        "[Postman KB2] Kiểm tra đối soát tồn kho P1 và ví C2 được bảo toàn tuyệt đối sau bù trừ",
        "Compensating Transaction đã khôi phục trạng thái kho chính xác: Không xảy ra hiện tượng thất thoát kho hay trừ tiền sai ví khi giao dịch bị gián đoạn."
    )

    add_heading_2("4.4. [KB2] Bước 4 – Kiểm tra trạng thái phiếu giữ hàng (Reservation) chuyển sang RELEASED")
    add_body("Gửi request GET http://localhost:8082/api/inventory/reservations để kiểm tra bản ghi giữ hàng của orderId=2 đã chuyển sang trạng thái RELEASED:")

    add_image_space(
        12,
        "[Postman KB2] Kiểm tra phiếu giữ hàng (Reservation) chuyển trạng thái RELEASED",
        "Xác nhận local transaction bồi hoàn tại MongoDB đã cập nhật bản ghi giữ hàng thành công từ RESERVED sang RELEASED."
    )

    add_heading_2("4.5. [KB2] Bước 5 – Kiểm tra chuỗi nhật ký bù trừ Saga Logs")
    add_body("Truy vấn lịch sử điều phối của đơn hàng số 2 qua GET http://localhost:8081/api/orders/2/saga-logs (ghi nhận SAGA FAILED, COMPENSATION SUCCESS, CANCEL_ORDER SUCCESS):")

    add_image_space(
        13,
        "[Postman KB2] Chuỗi nhật ký bù trừ Saga Logs của đơn hàng KB2 (FAILED -> RELEASE -> CANCEL)",
        "Saga Log chứng minh cơ chế bồi hoàn tự động: Khi Payment gặp lỗi, Orchestrator lập tức gọi release kho MongoDB và hủy đơn hàng MySQL."
    )

    add_heading_2("4.6. [KB2] Bước 6 – Xác minh dữ liệu trực tiếp trên 3 Database sau bù trừ")
    add_body("Truy vấn trực tiếp MySQL (orders id=2 CANCELLED, failure_reason 402), MongoDB (P1 stock=8, reservation RELEASED) và SQL Server (wallet C2 balance=100.0, không có payment):")

    add_image_space(
        14,
        "Kết quả truy vấn đối soát dữ liệu thực tế trên 3 Database sau bù trừ cho KB2",
        "Khẳng định hệ thống đạt tính nhất quán cuối cùng (Eventual Consistency): Toàn bộ các dịch vụ tham gia đều bảo toàn dữ liệu, không có bản ghi rác hay sai lệch số dư."
    )

    # =========================================================================
    # SECTION 5: COLLECTION RUNNER
    # =========================================================================
    add_heading_1("5. KIỂM THỬ TỰ ĐỘNG BẰNG POSTMAN COLLECTION RUNNER / NEWMAN CLI")
    add_body("Thực thi tự động toàn bộ test suite từ thư mục '0 - Kiểm tra service', 'KB1 - Happy path' đến 'KB2 - Thanh toán thất bại (402)' bằng Collection Runner:")

    add_image_space(
        15,
        "Kết quả chạy tự động toàn bộ Collection trên Postman Collection Runner hoàn thành 100% Passed",
        "Kết quả chạy tự động chứng minh toàn bộ các kịch bản kiểm thử Happy Path và Bù trừ phân tán hoạt động hoàn toàn chính xác, ổn định và đáng tin cậy."
    )

    # =========================================================================
    # SECTION 6: TỔNG KẾT VÀ BẢNG ĐỐI CHIẾU
    # =========================================================================
    add_heading_1("6. TỔNG KẾT VÀ ĐỐI CHIẾU KẾT QUẢ KIỂM THỬ SAGA PATTERN")

    add_heading_2("6.1. Bảng đối chiếu kết quả kiểm thử giữa 2 kịch bản trọng tâm")
    add_body("Bảng tổng hợp chi tiết diễn biến và kết quả của 2 kịch bản đã thực hiện kiểm thử:")

    headers = ["Thuộc tính / Bước kiểm thử", "Kịch bản 1 (KB1 – Happy Path)", "Kịch bản 2 (KB2 – Payment 402 Failed)"]
    rows_data = [
        ("Mã đơn hàng & Khách hàng", "Order ID: 1 | Customer ID: C1", "Order ID: 2 | Customer ID: C2"),
        ("Sản phẩm & Số lượng", "Product: P1 | Quantity: 2 (Đơn giá 1000/sp)", "Product: P1 | Quantity: 1 (Đơn giá 1000/sp)"),
        ("Tổng tiền & Số dư ví", "Tổng tiền: 2000.0 | Số dư C1: 5000.0 (Đủ tiền)", "Tổng tiền: 1000.0 | Số dư C2: 100.0 (Thiếu tiền)"),
        ("Bước 1: Tạo Order (orderdb)", "SUCCESS: Order status = PENDING", "SUCCESS: Order status = PENDING"),
        ("Bước 2: Giữ hàng (inventorydb)", "SUCCESS: MongoDB trừ tồn kho P1 (10 -> 8)", "SUCCESS: MongoDB trừ tồn kho P1 (8 -> 7)"),
        ("Bước 3: Thanh toán (paymentdb)", "SUCCESS: SQL Server trừ ví C1 (5000 -> 3000)", "FAILED: SQL Server trả về 402 Payment Required"),
        ("Bước bù trừ 1: Hoàn tác kho", "Không cần bù trừ (Happy Path thành công)", "SUCCESS: Release MongoDB trả kho P1 (7 -> 8)"),
        ("Bước bù trừ 2: Hủy đơn hàng", "Không cần thực hiện (Đơn được xác nhận)", "SUCCESS: Cập nhật MySQL Order thành CANCELLED"),
        ("Trạng thái đơn hàng cuối cùng", "Order: CONFIRMED | Saga: COMPLETED", "Order: CANCELLED | Saga: COMPENSATED"),
        ("Lý do lỗi (failureReason)", "null (Không phát sinh lỗi)", "402 - Số dư không đủ: còn 100.0, cần 1000.0"),
        ("Đối soát dữ liệu Polyglot", "P1: 8 | C1: 3000.0 | Reservation: RESERVED", "P1: 8 (Bảo toàn) | C2: 100.0 (Bảo toàn) | RELEASED")
    ]

    t_sum = doc.add_table(rows=len(rows_data) + 1, cols=3)
    t_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sum.autofit = False

    t_widths = [Inches(1.8), Inches(2.1), Inches(2.1)]
    for r in t_sum.rows:
        for idx, w in enumerate(t_widths):
            r.cells[idx].width = w

    b_spec_tbl = dict(val="single", sz="4", space="0", color="CBD5E1")
    for idx, h_text in enumerate(headers):
        cell = t_sum.rows[0].cells[idx]
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        set_cell_shading(cell, "1E3A8A")
        set_cell_border(cell, top=b_spec_tbl, bottom=b_spec_tbl, left=b_spec_tbl, right=b_spec_tbl)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(2)
        r = p.add_run(h_text)
        r.font.name, r.font.size, r.font.bold = "Times New Roman", Pt(10), True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row_values in enumerate(rows_data):
        row = t_sum.rows[r_idx + 1]
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            set_cell_shading(cell, bg)
            set_cell_border(cell, top=b_spec_tbl, bottom=b_spec_tbl, left=b_spec_tbl, right=b_spec_tbl)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(2)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                r.font.name, r.font.size, r.font.bold = "Times New Roman", Pt(9.5), True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                r.font.name, r.font.size = "Times New Roman", Pt(9.5)
                r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    sp_end = doc.add_paragraph()
    sp_end.style = 'Normal'
    sp_end.paragraph_format.space_before, sp_end.paragraph_format.space_after = Pt(6), Pt(6)

    add_heading_2("6.2. Đánh giá 3 tính chất cốt lõi của Compensating Transaction")
    add_body("Qua việc hiện thực và kiểm thử thực tế tại Kịch bản 2, hệ thống đã chứng minh đáp ứng trọn vẹn 3 yêu cầu lý thuyết bắt buộc của Giao dịch bù trừ (Compensating Transaction):")
    add_bullet("Hàm release trong inventory-service được thiết kế idempotent. Nếu gọi lặp lại nhiều lần với cùng một orderId, kết quả vẫn không làm thay đổi tồn kho lần thứ hai (chỉ hoàn kho nếu trạng thái đang là RESERVED).", "1. Idempotent (Tính bất biến khi gọi lại): ")
    add_bullet("Thao tác bù trừ hoàn tác trọn vẹn ảnh hưởng của thao tác thuận trước đó: Số lượng tồn kho được cộng bù chính xác đúng bằng số lượng đã giữ; reservation được đổi sang trạng thái RELEASED.", "2. Reversible (Tính khả nghịch): ")
    add_bullet("Compensating Transaction được bọc trong RetryTemplate với cơ chế thử lại tối đa 3 lần. Nếu xảy ra lỗi mạng tạm thời, Orchestrator sẽ kiên trì retry nhằm đảm bảo bù trừ thành công, không để hệ thống rơi vào trạng thái dữ liệu rác.", "3. Reliable (Tính tin cậy cao): ")

    # Save document
    doc.save(output_docx)
    print(f"Document successfully created at: {output_docx} (Size: {os.path.getsize(output_docx) // 1024} KB)")

if __name__ == '__main__':
    build_slot13_docx()
